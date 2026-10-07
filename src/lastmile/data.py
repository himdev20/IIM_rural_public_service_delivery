"""Data loading.

The routing network is represented by a single :class:`Instance`:

* node ``0`` is the depot (the Frozen Serum Laboratory, FSL);
* nodes ``1..n`` are livestock locations, numbered exactly as in the case;
* ``dist`` is a symmetric ``(n+1, n+1)`` matrix in km;
* ``demand[i]`` is the number of 50-litre cryocans location ``i`` needs per visit.

Node ids therefore coincide with matrix indices everywhere in the code base.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field, replace
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

DEPOT = 0


@dataclass(frozen=True)
class Instance:
    """A capacitated routing network with a single depot at node 0."""

    name: str
    dist: np.ndarray
    demand: np.ndarray
    labels: tuple[str, ...]
    #: cells (i, j) that the source workbook stores as *text* rather than numbers.
    #: Only needed to replicate Excel's sort order in the VBA model exactly.
    text_cells: frozenset[tuple[int, int]] = field(default_factory=frozenset)

    def __post_init__(self) -> None:
        d = np.asarray(self.dist, dtype=float)
        if d.ndim != 2 or d.shape[0] != d.shape[1]:
            raise ValueError("dist must be a square matrix")
        if not np.allclose(d, d.T):
            raise ValueError("dist must be symmetric")
        if len(self.demand) != d.shape[0] or len(self.labels) != d.shape[0]:
            raise ValueError("demand/labels must have one entry per node incl. depot")
        object.__setattr__(self, "dist", d)
        object.__setattr__(self, "demand", np.asarray(self.demand, dtype=int))

    # ------------------------------------------------------------------ basics
    @property
    def n(self) -> int:
        """Number of customer locations (excluding the depot)."""
        return self.dist.shape[0] - 1

    @property
    def customers(self) -> tuple[int, ...]:
        return tuple(range(1, self.n + 1))

    @property
    def depot_dist(self) -> np.ndarray:
        """Distance from the depot to every node (index 0 is 0)."""
        return self.dist[DEPOT]

    def fingerprint(self) -> str:
        """Stable short hash of the data, used as a cache key (memoised)."""
        cached = self.__dict__.get("_fingerprint")
        if cached is None:
            h = hashlib.sha1()
            h.update(np.round(self.dist, 6).tobytes())
            h.update(self.demand.astype(np.int64).tobytes())
            cached = h.hexdigest()[:12]
            object.__setattr__(self, "_fingerprint", cached)
        return cached

    # -------------------------------------------------------------- variants
    def with_demand(self, demand: np.ndarray, name: str | None = None) -> "Instance":
        return replace(self, demand=np.asarray(demand, dtype=int), name=name or self.name)

    def metric_closure(self) -> "Instance":
        """Replace every distance by the shortest-path distance (Floyd-Warshall).

        The supplied matrix violates the triangle inequality for many pairs; a
        driver could take the shorter path through a third location. This is a
        robustness variant, not the replication matrix.
        """
        d = self.dist.copy()
        for k in range(d.shape[0]):
            d = np.minimum(d, d[:, [k]] + d[[k], :])
        return replace(self, dist=d, name=f"{self.name}-metric", text_cells=frozenset())

    def triangle_violation_share(self) -> float:
        """Share of node pairs whose direct distance exceeds the shortest path."""
        closed = self.metric_closure().dist
        iu = np.triu_indices_from(self.dist, k=1)
        return float(np.mean(self.dist[iu] - closed[iu] > 1e-9))


# ---------------------------------------------------------------------------
# Case data (RPAHD workbook)
# ---------------------------------------------------------------------------
def load_case_instance(workbook: Path) -> Instance:
    """Load the 177-location Nagpur instance from the RPAHD workbook.

    Sheets used: ``Dmatrix`` (177x177 inter-location km) and ``Demands``
    (distance from FSL, demand, location name).
    """
    wb = load_workbook(workbook, data_only=True, read_only=True)
    ws = wb["Dmatrix"]
    raw_rows = list(ws.iter_rows(min_row=2, min_col=2, values_only=True))
    raw_rows = [r for r in raw_rows if r and r[0] is not None]
    n = len(raw_rows)
    text_cells: set[tuple[int, int]] = set()
    inter = np.zeros((n, n))
    for i, row in enumerate(raw_rows):
        for j, value in enumerate(row[:n]):
            if isinstance(value, str):
                text_cells.add((i + 1, j + 1))
            inter[i, j] = float(value)

    demands = list(wb["Demands"].iter_rows(min_row=2, max_row=n + 1, max_col=5, values_only=True))
    depot_km = np.array([float(r[1]) for r in demands])
    demand = np.array([int(r[2]) for r in demands])
    names = tuple(str(r[4]).strip() for r in demands)
    wb.close()

    dist = np.zeros((n + 1, n + 1))
    dist[1:, 1:] = inter
    dist[0, 1:] = depot_km
    dist[1:, 0] = depot_km
    return Instance(
        name="nagpur",
        dist=dist,
        demand=np.concatenate([[0], demand]),
        labels=("FSL",) + names,
        text_cells=frozenset(text_cells),
    )


@dataclass(frozen=True)
class StatusQuoRoute:
    number: int
    name: str
    sequence: tuple[int, ...]
    reported_km: float


def load_status_quo(workbook: Path) -> list[StatusQuoRoute]:
    """Load Table 1 (the nine judgment-based routes) from the case workbook."""
    wb = load_workbook(workbook, data_only=True, read_only=True)
    routes = []
    for row in wb["Table 1"].iter_rows(min_row=1, values_only=True):
        if isinstance(row[0], (int, float)) and isinstance(row[2], str) and "-" in row[2]:
            seq = tuple(int(x) for x in row[2].split("-"))
            routes.append(StatusQuoRoute(int(row[0]), str(row[1]).strip(), seq, float(row[4])))
    wb.close()
    return routes


def load_spreadsheet_solution(model_file: Path, sheet: str) -> dict:
    """Read a solution sheet (e.g. ``Sheet20``) produced by the VBA savings model."""
    wb = load_workbook(model_file, data_only=True, read_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    total = float(rows[0][9])
    routes, route_km = [], []
    for row in rows:
        if len(row) > 5 and row[1] == "Route Distance:":
            routes.append([int(v) for v in row[5:] if v is not None])
            route_km.append(float(row[3]))
    return {"total_km": total, "routes": routes, "route_km": route_km}


def load_spreadsheet_sort(model_file: Path) -> list[tuple[int, int]]:
    """The sorted savings list (pairs i>j) stored in the workbook's ``Sort`` sheet."""
    wb = load_workbook(model_file, data_only=True, read_only=True)
    pairs = [
        (int(r[1]), int(r[2]))
        for r in wb["Sort"].iter_rows(min_row=2, max_col=4, values_only=True)
        if r[0] is not None
    ]
    wb.close()
    return pairs


# ---------------------------------------------------------------------------
# Small teaching instance (ABC warehouse, 13 PDS shops) used in tests
# ---------------------------------------------------------------------------
def load_abc_instance(csv_file: Path) -> Instance:
    df = pd.read_csv(csv_file)
    n = len(df)
    dist = np.zeros((n + 1, n + 1))
    dist[1:, 1:] = df[[f"C{i}" for i in range(1, n + 1)]].to_numpy(float)
    dist[0, 1:] = dist[1:, 0] = df["depot_km"].to_numpy(float)
    return Instance(
        name="abc-warehouse",
        dist=dist,
        demand=np.concatenate([[0], df["demand"].to_numpy(int)]),
        labels=("ABC",) + tuple(df["node"]),
    )


def savings_matrix(inst: Instance) -> np.ndarray:
    """Clarke-Wright savings s(i, j) = d(0, i) + d(0, j) - d(i, j); zero diagonal."""
    d0 = inst.depot_dist
    s = d0[:, None] + d0[None, :] - inst.dist
    s[0, :] = s[:, 0] = 0.0
    np.fill_diagonal(s, 0.0)
    return s
